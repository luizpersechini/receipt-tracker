"""Build the front-page expense table as an editable .xlsx — the spreadsheet
twin of the PDF front page.

Mirrors the Atlantic Nickel "Prestação de Contas - Cartão Corporativo" form
(MMB-F-REG-0101-001-00): header band (logo · title · código · página), the
identity band (Nome / Vencimento / Nº Cartão), the categorised expense table
with per-column SUM totals, and a signature line. No receipt images — that's
the PDF's job.

Reuses `report_rows_view()` so the numbers/labels match the PDF exactly.
"""

from __future__ import annotations

import io
import logging
import os
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from ..config import settings
from .expense_report_pdf import report_rows_view

logger = logging.getLogger(__name__)

_ARIAL = "Arial"
_THIN = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
_RIGHT = Alignment(horizontal="right", vertical="center")
_NUM_FMT = "#,##0.00"
_BRL_FMT = '"R$"\\ #,##0.00'

_LEAD_COLS = 4  # Data | Centro de Custo+Uneg | Conta Contábil | Detalhamento


def _logo_path() -> Optional[str]:
    for c in (
        settings.data_dir.parent / "static" / "report_assets" / "logo.png",
        Path("static/report_assets/logo.png"),
    ):
        if os.path.exists(c):
            return str(c)
    return None


def build_statement_xlsx(
    statement_id: int,
    user_id: Optional[int] = None,
    vencimento: Optional[str] = None,
) -> bytes:
    """Render the statement's front-page expense table as .xlsx bytes."""
    view = report_rows_view(statement_id, user_id, vencimento)
    cols = view["columns"]  # [{key, label}]
    ncat = len(cols)
    total_col = _LEAD_COLS + ncat + 1  # 1-based index of the TOTAL column
    cat_col = {c["key"]: _LEAD_COLS + 1 + i for i, c in enumerate(cols)}

    title_font = Font(name=_ARIAL, size=16, bold=True)
    hdr_font = Font(name=_ARIAL, size=10, bold=True)
    cell_font = Font(name=_ARIAL, size=10)
    small_font = Font(name=_ARIAL, size=8)

    wb = Workbook()
    ws = wb.active
    ws.title = "Prestação de Contas"

    widths = [15.5, 18.0, 14.2, 48.7] + [16.5] * ncat + [15.8]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    def put(row, col, value, *, font=cell_font, align=_LEFT, fmt=None, border=False):
        c = ws.cell(row=row, column=col, value=value)
        c.font = font
        c.alignment = align
        if fmt:
            c.number_format = fmt
        if border:
            c.border = _BORDER
        return c

    # ── Header band (rows 1-4) ────────────────────────────────────────────
    ws.merge_cells(start_row=1, start_column=1, end_row=4, end_column=3)  # logo box
    ws.merge_cells(start_row=1, start_column=4, end_row=4, end_column=total_col - 2)
    put(1, 4, view.get("report_title") or "", font=title_font, align=_CENTER)
    gcol = total_col - 1
    put(1, gcol, "Código", font=small_font, align=_CENTER)
    ws.merge_cells(start_row=1, start_column=total_col, end_row=2, end_column=total_col)
    put(1, total_col, view.get("codigo") or "", font=small_font, align=_CENTER)
    put(3, gcol, "Página", font=small_font, align=_CENTER)
    ws.merge_cells(start_row=3, start_column=total_col, end_row=4, end_column=total_col)
    put(3, total_col, "01 / 01", font=small_font, align=_CENTER)

    lp = _logo_path()
    if lp:
        try:
            from openpyxl.drawing.image import Image as XLImage

            img = XLImage(lp)
            img.width, img.height = 200, 58
            ws.add_image(img, "A1")
        except Exception as e:  # pragma: no cover — image libs/format issues
            logger.warning("xlsx logo embed failed: %s", e)

    # ── Identity band (rows 5-7) ──────────────────────────────────────────
    def band(row, label, value):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        put(row, 1, label, font=hdr_font, align=_LEFT)
        ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=total_col)
        put(row, 4, value, font=cell_font, align=_LEFT)

    band(5, "Nome:", view.get("nome") or "")
    band(6, "VENCIMENTO:", view.get("vencimento") or "")
    band(7, "Nº Cartão:", view.get("num_cartao") or "")

    # ── Table header (row 8) ──────────────────────────────────────────────
    hdr_row = 8
    headers = [
        "Data",
        "Centro de Custo+Uneg",
        "Conta Contábil",
        "Detalhamento da despesa",
        *[c["label"] for c in cols],
        "TOTAL",
    ]
    for j, h in enumerate(headers, start=1):
        put(hdr_row, j, h, font=hdr_font, align=_CENTER, border=True)
    ws.row_dimensions[hdr_row].height = 28

    # ── Data rows ─────────────────────────────────────────────────────────
    r = hdr_row + 1
    first_data = r
    for row in view["rows"]:
        put(r, 1, row["date"], align=_CENTER, border=True)
        put(r, 2, row.get("cost_center_effective") or "", align=_CENTER, border=True)
        put(r, 3, row.get("conta") or "", align=_CENTER, border=True)
        put(r, 4, row.get("detalhe") or "", align=_LEFT, border=True)
        for c in cols:
            col = cat_col[c["key"]]
            val = row["amount"] if c["key"] == row["col"] else None
            put(r, col, val, align=_RIGHT, fmt=_NUM_FMT, border=True)
        put(r, total_col, row["amount"], align=_RIGHT, fmt=_NUM_FMT, border=True)
        r += 1

    if view.get("iof"):
        outros_key = cols[-1]["key"]
        put(r, 1, "", border=True)
        put(r, 2, "", border=True)
        put(r, 3, "", border=True)
        put(r, 4, "IOF", align=_LEFT, border=True)
        for c in cols:
            col = cat_col[c["key"]]
            val = view["iof"] if c["key"] == outros_key else None
            put(r, col, val, align=_RIGHT, fmt=_NUM_FMT, border=True)
        put(r, total_col, view["iof"], align=_RIGHT, fmt=_NUM_FMT, border=True)
        r += 1
    last_data = r - 1

    # ── Totals row (live SUM formulas) ────────────────────────────────────
    put(r, 1, "", border=True)
    put(r, 2, "", border=True)
    put(r, 3, "", border=True)
    put(r, 4, "TOTAL", font=hdr_font, align=_RIGHT, border=True)
    for c in cols:
        col = cat_col[c["key"]]
        letter = get_column_letter(col)
        put(
            r,
            col,
            f"=SUM({letter}{first_data}:{letter}{last_data})",
            font=hdr_font,
            align=_RIGHT,
            fmt=_BRL_FMT,
            border=True,
        )
    tl = get_column_letter(total_col)
    put(
        r,
        total_col,
        f"=SUM({tl}{first_data}:{tl}{last_data})",
        font=hdr_font,
        align=_RIGHT,
        fmt=_BRL_FMT,
        border=True,
    )

    # ── Signature line ────────────────────────────────────────────────────
    sig_row = r + 3
    put(sig_row, 1, "Data", font=hdr_font, align=_LEFT)
    ws.merge_cells(start_row=sig_row, start_column=4, end_row=sig_row, end_column=total_col)
    put(
        sig_row,
        4,
        f"Assinatura {view.get('nome') or ''} + Aprovador imediato:",
        font=hdr_font,
        align=_LEFT,
    )

    buf = io.BytesIO()
    wb.save(buf)
    logger.info(
        "Built XLSX report for statement %d: %d rows, %d bytes",
        statement_id,
        len(view["rows"]),
        buf.tell(),
    )
    return buf.getvalue()
