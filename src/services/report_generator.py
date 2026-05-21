"""Excel report generator."""

import json
from datetime import datetime
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..config import settings
from .fx_fetcher import convert
from .matching_engine import match_receipts_to_statements


HEADER_FILL = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
ALT_FILL = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
GREEN_FILL = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
RED_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
SUBTOTAL_FILL = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")
TOTAL_FILL = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")
TOTAL_FONT = Font(bold=True, color="FFFFFF", size=11)
THIN_BORDER = Border(
    bottom=Side(border_style="thin", color="E4E6EA"),
    top=Side(border_style="thin", color="E4E6EA"),
)


def _auto_width(ws):
    for col in ws.columns:
        max_len = 0
        letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if len(str(cell.value or "")) > max_len:
                    max_len = len(str(cell.value))
            except Exception:
                pass
        ws.column_dimensions[letter].width = min(max_len + 3, 40)


def _write_header(ws, row: int, headers: list):
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _parse_meta(receipt: dict) -> dict:
    meta = receipt.get("metadata")
    if isinstance(meta, str):
        try:
            return json.loads(meta)
        except json.JSONDecodeError:
            return {}
    return meta or {}


def _fmt_amount(val) -> Optional[float]:
    if val is None:
        return None
    try:
        return float(val)
    except (TypeError, ValueError):
        return None


def _converted_label(target_currency: str) -> str:
    return f"Converted ({target_currency})"


def _write_receipt_rows(ws, receipts: list, start_row: int,
                        target_currency: str, fx_rates: dict,
                        show_expense_type: bool = False,
                        show_match_status: bool = False) -> int:
    row = start_row
    for i, r in enumerate(receipts):
        meta = _parse_meta(r)
        fill = ALT_FILL if i % 2 == 0 else None

        amount = _fmt_amount(meta.get("total_amount"))
        currency = (meta.get("currency") or "").upper()
        converted = convert(amount, currency, target_currency, fx_rates) if currency else None

        col = 1
        cells = [
            r.get("id"),
            meta.get("date"),
            meta.get("vendor_name"),
            meta.get("city") or meta.get("vendor_address", "")[:30] if meta.get("vendor_address") else meta.get("city"),
            amount,
            currency or None,
            converted,
            meta.get("category"),
            r.get("description"),
        ]
        if show_expense_type:
            cells.insert(8, r.get("expense_type", "corp_card"))
        if show_match_status:
            cells.append(r.get("match_status", "pending"))

        for val in cells:
            c = ws.cell(row=row, column=col, value=val)
            if fill:
                c.fill = fill
            c.border = THIN_BORDER
            col += 1
        row += 1

    return row


def _receipt_headers(target_currency: str, show_expense_type=False, show_match_status=False):
    h = ["ID", "Date", "Vendor", "Location", "Amount", "Currency",
         _converted_label(target_currency), "Category", "Description"]
    if show_expense_type:
        h.insert(8, "Type")
    if show_match_status:
        h.append("Match Status")
    return h


def _write_category_subtotals(ws, receipts: list, row: int,
                              target_currency: str, fx_rates: dict) -> int:
    totals: dict = {}
    for r in receipts:
        meta = _parse_meta(r)
        cat = meta.get("category") or "other"
        amount = _fmt_amount(meta.get("total_amount")) or 0.0
        currency = (meta.get("currency") or "BRL").upper()
        converted = convert(amount, currency, target_currency, fx_rates) or 0.0
        totals[cat] = totals.get(cat, 0.0) + converted

    ws.cell(row=row, column=1, value="SUBTOTALS BY CATEGORY").font = Font(bold=True, size=10)
    row += 1
    for cat, total in sorted(totals.items()):
        c1 = ws.cell(row=row, column=1, value=cat.replace("_", " ").title())
        c2 = ws.cell(row=row, column=2, value=round(total, 2))
        c3 = ws.cell(row=row, column=3, value=target_currency)
        for c in (c1, c2, c3):
            c.fill = SUBTOTAL_FILL
            c.font = Font(italic=True, size=10)
        row += 1

    grand = sum(totals.values())
    gt1 = ws.cell(row=row, column=1, value="GRAND TOTAL")
    gt2 = ws.cell(row=row, column=2, value=round(grand, 2))
    gt3 = ws.cell(row=row, column=3, value=target_currency)
    for c in (gt1, gt2, gt3):
        c.fill = TOTAL_FILL
        c.font = TOTAL_FONT
    row += 2
    return row


def generate_excel_report(receipts: list, statements: list,
                          target_currency: str = "BRL",
                          fx_rates: Optional[dict] = None) -> str:
    if fx_rates is None:
        fx_rates = {}

    settings.reports_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    report_path = settings.reports_dir / f"receipt_report_{timestamp}.xlsx"

    wb = Workbook()

    ws_summary = wb.active
    ws_summary.title = "Summary"
    _create_summary_sheet(ws_summary, receipts, statements, target_currency, fx_rates)

    ws_receipts = wb.create_sheet("Receipts")
    _create_receipts_sheet(ws_receipts, receipts, target_currency, fx_rates)

    ws_statements = wb.create_sheet("Statements")
    _create_statements_sheet(ws_statements, statements)

    ws_recon = wb.create_sheet("Reconciliation")
    _create_reconciliation_sheet(ws_recon, receipts, statements, target_currency, fx_rates)

    wb.save(str(report_path))
    return str(report_path)


def _create_summary_sheet(ws, receipts, statements, target_currency, fx_rates):
    ws.merge_cells("A1:E1")
    ws["A1"] = "Receipt Tracker — Treasury Report"
    ws["A1"].font = Font(size=16, bold=True, color="1E1B4B")
    ws["A1"].alignment = Alignment(horizontal="center")

    data = [
        ("Report Generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Target Currency", target_currency),
        ("Total Receipts", len(receipts)),
        ("Total Statements", len(statements)),
    ]

    total_original: dict = {}
    total_converted = 0.0
    for r in receipts:
        meta = _parse_meta(r)
        amt = _fmt_amount(meta.get("total_amount")) or 0.0
        cur = (meta.get("currency") or "BRL").upper()
        total_original[cur] = total_original.get(cur, 0.0) + amt
        cv = convert(amt, cur, target_currency, fx_rates)
        if cv is not None:
            total_converted += cv

    for cur, amt in sorted(total_original.items()):
        data.append((f"Total ({cur})", round(amt, 2)))
    data.append((f"Grand Total ({target_currency})", round(total_converted, 2)))

    for i, (label, value) in enumerate(data, 3):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        ws.cell(row=i, column=2, value=value)

    _auto_width(ws)


def _create_receipts_sheet(ws, receipts, target_currency, fx_rates):
    headers = _receipt_headers(target_currency, show_expense_type=True, show_match_status=True)
    _write_header(ws, 1, headers)
    next_row = _write_receipt_rows(ws, receipts, 2, target_currency, fx_rates,
                                   show_expense_type=True, show_match_status=True)
    _write_category_subtotals(ws, receipts, next_row + 1, target_currency, fx_rates)
    _auto_width(ws)


def _create_statements_sheet(ws, statements):
    headers = ["Statement ID", "Filename", "Transaction Date",
               "Description", "Amount", "Category"]
    _write_header(ws, 1, headers)
    row = 2
    for s in statements:
        for txn in s.get("transactions", []):
            cells = [s.get("id"), s.get("filename"), txn.get("date"),
                     txn.get("description"), txn.get("amount"), txn.get("category")]
            for col, val in enumerate(cells, 1):
                ws.cell(row=row, column=col, value=val).border = THIN_BORDER
            row += 1
    _auto_width(ws)


def _create_reconciliation_sheet(ws, receipts, statements, target_currency, fx_rates):
    headers = [
        "Receipt ID", "Receipt Date", "Receipt Vendor", "Receipt Amount", "Currency",
        _converted_label(target_currency),
        "Statement ID", "Txn Date", "Txn Description", "Txn Amount",
        "Match Score", "Status",
    ]
    _write_header(ws, 1, headers)

    result = match_receipts_to_statements(receipts, statements)
    row = 2
    for m in result["matches"]:
        cur = (m.get("receipt_currency") or "BRL").upper()
        amt = _fmt_amount(m.get("receipt_amount"))
        cv = convert(amt, cur, target_currency, fx_rates)
        score = m.get("match_score", 0)
        fill = GREEN_FILL if score >= 0.8 else (YELLOW_FILL if score >= 0.6 else RED_FILL)
        cells = [
            m.get("receipt_id"), m.get("receipt_date"), m.get("receipt_vendor"),
            amt, cur, cv,
            m.get("statement_id"), m.get("transaction_date"),
            m.get("transaction_description"), m.get("transaction_amount"),
            score,
        ]
        for col, val in enumerate(cells, 1):
            ws.cell(row=row, column=col, value=val).border = THIN_BORDER
        status_cell = ws.cell(row=row, column=len(cells) + 1,
                              value="MATCHED" if score >= 0.8 else "PARTIAL" if score >= 0.6 else "WEAK")
        status_cell.fill = fill
        row += 1
    _auto_width(ws)


def generate_expense_report(receipts: list, report_type: str,
                            target_currency: str = "BRL",
                            fx_rates: Optional[dict] = None,
                            matches: Optional[list] = None) -> str:
    """Generate a typed expense report Excel file. report_type: corp_card|reimbursement|all."""
    if fx_rates is None:
        fx_rates = {}

    settings.reports_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    label = report_type.replace("_", "-")
    report_path = settings.reports_dir / f"expense_report_{label}_{timestamp}.xlsx"

    wb = Workbook()

    ws_summary = wb.active
    ws_summary.title = "Summary"
    _create_expense_summary(ws_summary, receipts, report_type, target_currency, fx_rates)

    if report_type in ("corp_card", "all"):
        corp = [r for r in receipts if r.get("expense_type") == "corp_card"]
        ws_corp = wb.create_sheet("Corp Card")
        _create_corp_card_sheet(ws_corp, corp, target_currency, fx_rates, matches or [])

    if report_type in ("reimbursement", "all"):
        reimb = [r for r in receipts if r.get("expense_type") == "reimbursement"]
        ws_reimb = wb.create_sheet("Reimbursement")
        _create_reimbursement_sheet(ws_reimb, reimb, target_currency, fx_rates)

    wb.save(str(report_path))
    return str(report_path)


def _create_expense_summary(ws, receipts, report_type, target_currency, fx_rates):
    ws.merge_cells("A1:F1")
    ws["A1"] = f"Expense Report — {report_type.replace('_', ' ').title()}"
    ws["A1"].font = Font(size=16, bold=True, color="1E1B4B")
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.cell(row=2, column=1, value="Generated").font = Font(bold=True)
    ws.cell(row=2, column=2, value=datetime.now().strftime("%Y-%m-%d %H:%M"))
    ws.cell(row=3, column=1, value="Report Type").font = Font(bold=True)
    ws.cell(row=3, column=2, value=report_type.replace("_", " ").title())
    ws.cell(row=4, column=1, value="Target Currency").font = Font(bold=True)
    ws.cell(row=4, column=2, value=target_currency)

    types = {"corp_card": "Corp Card", "reimbursement": "Reimbursement"}
    row = 6
    ws.cell(row=row, column=1, value="Breakdown by Type").font = Font(bold=True, size=11)
    row += 1
    _write_header(ws, row, ["Type", "Receipts", f"Total ({target_currency})"])
    row += 1

    for etype, label in types.items():
        subset = [r for r in receipts if r.get("expense_type") == etype]
        if not subset:
            continue
        total_cv = sum(
            (convert(
                _fmt_amount(_parse_meta(r).get("total_amount")) or 0,
                (_parse_meta(r).get("currency") or "BRL").upper(),
                target_currency, fx_rates
            ) or 0)
            for r in subset
        )
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=len(subset))
        ws.cell(row=row, column=3, value=round(total_cv, 2))
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Breakdown by Category").font = Font(bold=True, size=11)
    row += 1
    _write_header(ws, row, ["Category", f"Total ({target_currency})"])
    row += 1

    cat_totals: dict = {}
    for r in receipts:
        meta = _parse_meta(r)
        cat = (meta.get("category") or "other").replace("_", " ").title()
        amt = _fmt_amount(meta.get("total_amount")) or 0
        cur = (meta.get("currency") or "BRL").upper()
        cv = convert(amt, cur, target_currency, fx_rates) or 0
        cat_totals[cat] = cat_totals.get(cat, 0.0) + cv

    for cat, total in sorted(cat_totals.items()):
        ws.cell(row=row, column=1, value=cat)
        c = ws.cell(row=row, column=2, value=round(total, 2))
        c.fill = SUBTOTAL_FILL
        row += 1

    grand = sum(cat_totals.values())
    ws.cell(row=row, column=1, value="GRAND TOTAL").font = TOTAL_FONT
    ws.cell(row=row, column=1).fill = TOTAL_FILL
    c = ws.cell(row=row, column=2, value=round(grand, 2))
    c.font = TOTAL_FONT
    c.fill = TOTAL_FILL

    _auto_width(ws)


def _create_corp_card_sheet(ws, receipts, target_currency, fx_rates, matches):
    """Corp card sheet: receipt rows + matched statement transaction column."""
    match_by_receipt = {m["receipt_id"]: m for m in matches}

    headers = _receipt_headers(target_currency, show_match_status=True) + \
              ["Txn Date", "Txn Description", "Txn Amount", "Match Score"]
    _write_header(ws, 1, headers)

    row = 2
    for i, r in enumerate(receipts):
        meta = _parse_meta(r)
        amount = _fmt_amount(meta.get("total_amount"))
        currency = (meta.get("currency") or "BRL").upper()
        converted = convert(amount, currency, target_currency, fx_rates)
        fill = ALT_FILL if i % 2 == 0 else None

        m = match_by_receipt.get(r["id"])
        base = [
            r.get("id"), meta.get("date"), meta.get("vendor_name"),
            meta.get("city"), amount, currency, converted,
            meta.get("category"), r.get("description"),
            r.get("match_status", "pending"),
            m.get("transaction_date") if m else None,
            m.get("transaction_description") if m else None,
            m.get("transaction_amount") if m else None,
            round(m.get("match_score", 0), 2) if m else None,
        ]
        for col, val in enumerate(base, 1):
            c = ws.cell(row=row, column=col, value=val)
            if fill:
                c.fill = fill
            c.border = THIN_BORDER
        row += 1

    _write_category_subtotals(ws, receipts, row + 1, target_currency, fx_rates)
    _auto_width(ws)


def _create_reimbursement_sheet(ws, receipts, target_currency, fx_rates):
    headers = _receipt_headers(target_currency)
    _write_header(ws, 1, headers)
    next_row = _write_receipt_rows(ws, receipts, 2, target_currency, fx_rates)
    _write_category_subtotals(ws, receipts, next_row + 1, target_currency, fx_rates)
    _auto_width(ws)
